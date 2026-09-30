"""Parse the reviewed workbook layout; reject changes instead of guessing columns."""
from datetime import datetime
import math
import re
import openpyxl
from common import CODES, PREFECTURES, day, require, timestamp

SHEET = '公表資料（都道府県別）'
AGGREGATES = {'東北局', '関東局', '中部局', '近畿局', '中国局', '四国局', '九州局', '九州沖縄局', '全国'}
NAME_CODES = dict(zip(PREFECTURES, CODES)) | {'北海道局': 'JP-01', '沖縄局': 'JP-47'}

def compact(value):
    return re.sub(r'\s+', '', str(value or ''))

def date_cell(value):
    require(isinstance(value, datetime), 'Workbook date cell must be an Excel date')
    return value.date().isoformat()

def validate_prices(data, previous=None):
    require(data.get('schemaVersion') == 1 and data.get('unit') == 'JPY/L' and data.get('basis') == 'CASH_TAX_INCLUDED_PREFECTURAL_REFERENCE', 'Invalid price schema/units')
    survey, published = day(data['surveyDate']), day(data['publishedAt'])
    require(survey <= published <= timestamp(data['fetchedAt']).date(), 'Price dates inconsistent')
    require(data.get('source') == 'meti-prices' and data.get('sourceUrl', '').startswith('https://www.enecho.meti.go.jp/statistics/petroleum_and_lpgas/pl007/xlsx/'), 'Invalid price source')
    rows = data.get('records', [])
    require(len(rows) == 141, 'Expected exactly 47 x 3 = 141 price records')
    keys = set()
    for row in rows:
        key = (row.get('prefectureCode'), row.get('fuelType'))
        require(key[0] in CODES and key[1] in {'REGULAR', 'HIGH_OCTANE', 'DIESEL'} and key not in keys, 'Duplicate/invalid prefecture/fuel')
        keys.add(key)
        value = row.get('priceJpy')
        require(type(value) in {int, float} and math.isfinite(value) and 50 <= value <= 400, 'Price outside 50..400 JPY/L')
        for field in ['surveyDate', 'publishedAt', 'fetchedAt', 'sourceUrl']:
            require(row.get(field) == data[field], f'Price record {field} mismatch')
    if previous:
        require(survey >= day(previous['surveyDate']) and published >= day(previous['publishedAt']), 'Older price replacement blocked')

def parse_prices(path, metadata):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        require(SHEET in workbook.sheetnames, 'Official prefectural sheet missing')
        sheet = workbook[SHEET]
        require(compact(sheet['B1'].value) == '石油製品小売市況調査(都道府県別)', 'Unexpected table title')
        require(compact(sheet['B3'].value) == '現金価格（消費税込み）', 'Unexpected tax/cash basis')
        for cell, expected in {'C5': 'ハイオク（\\/㍑）', 'E5': 'レギュラー（\\/㍑）', 'G5': '軽油店頭（\\/㍑）'}.items():
            require(compact(sheet[cell].value) == expected, f'Unexpected fuel/unit header: {cell}')
        survey = date_cell(sheet['D6'].value)
        require(all(date_cell(sheet[c].value) == survey for c in ['F6', 'H6']), 'Latest survey dates differ')
        require(all(date_cell(sheet[c].value) < survey for c in ['C6', 'E6', 'G6']), 'Previous survey columns invalid')
        published = date_cell(sheet['E2'].value)
        require(metadata.get('publishedAt') == published and metadata.get('surveyDate') == survey, 'Source proof dates differ from workbook')
        data = {'schemaVersion': 1, 'source': 'meti-prices', 'unit': 'JPY/L', 'basis': 'CASH_TAX_INCLUDED_PREFECTURAL_REFERENCE', 'surveyDate': survey, 'publishedAt': published, 'fetchedAt': metadata['fetchedAt'], 'sourceUrl': metadata['sourceUrl'], 'records': []}
        seen = set()
        for row in sheet.iter_rows(min_row=7, values_only=True):
            name = compact(row[1])
            if name in AGGREGATES or not name:
                continue
            require(name in NAME_CODES, f'Unrecognized prefecture/aggregate: {name}')
            code = NAME_CODES[name]
            require(code not in seen, f'Duplicate prefecture: {code}')
            seen.add(code)
            for index, fuel in [(3, 'HIGH_OCTANE'), (5, 'REGULAR'), (7, 'DIESEL')]:
                data['records'].append({'prefectureCode': code, 'fuelType': fuel, 'priceJpy': row[index], **{k: data[k] for k in ['surveyDate', 'publishedAt', 'fetchedAt', 'sourceUrl']}})
        require(seen == set(CODES), 'Missing prefecture coverage')
        data['records'].sort(key=lambda r: (r['prefectureCode'], r['fuelType']))
        validate_prices(data)
        return data
    finally:
        workbook.close()
