import copy
from datetime import datetime, timedelta, timezone
import io
import json
from pathlib import Path
import shutil
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from urllib.error import HTTPError

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts/importer'))
import openpyxl
from common import CODES, PREFECTURES, count_gate, digest, encode, verified_input
from osm import parse_osm, tri, station
from prices import parse_prices, validate_prices, SHEET
from pipeline import update, fetch_price


def workbook(path):
    w = openpyxl.Workbook()
    s = w.active
    s.title = SHEET
    for cell, value in {'B1': '石油製品小売市況調査(都道府県別)', 'B3': '現金価格（消費税込み）', 'C5': 'ハイオク（\\/㍑）', 'E5': 'レギュラー（\\/㍑）', 'G5': '軽 油 店 頭（\\/㍑）'}.items():
        s[cell] = value
    s['E2'] = datetime(2026, 9, 30)
    for cell in ['D6', 'F6', 'H6']:
        s[cell] = datetime(2026, 9, 28)
    for cell in ['C6', 'E6', 'G6']:
        s[cell] = datetime(2026, 9, 14)
    for row, name in enumerate(PREFECTURES, 7):
        s.cell(row, 2, name + '局' if name in {'北海道', '沖縄'} else name)
        for col, value in [(4, 181.1), (6, 170.2), (8, 159.3), (10, 2538)]:
            s.cell(row, col, value)
    w.save(path)
    return w


def metadata(path):
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        published, survey = book[SHEET]['E2'].value, book[SHEET]['D6'].value
    finally:
        book.close()
    return {'sourceUrl': f'https://www.enecho.meti.go.jp/statistics/petroleum_and_lpgas/pl007/xlsx/{published:%y%m%d}.xlsx', 'fetchedAt': datetime.now(timezone.utc).isoformat(), 'sha256': digest(path), 'bytes': path.stat().st_size, 'surveyDate': survey.date().isoformat(), 'publishedAt': published.date().isoformat()}


class PriceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'input.xlsx'
        self.book = workbook(self.path)
        self.addCleanup(self.book.close)
        self.meta = metadata(self.path)

    def parse(self):
        return parse_prices(self.path, self.meta)

    def test_exact_47_by_3_units_dates_and_ignore_kerosene(self):
        data = self.parse()
        self.assertEqual(len(data['records']), 141)
        self.assertEqual({r['prefectureCode'] for r in data['records']}, set(CODES))
        self.assertEqual(data['unit'], 'JPY/L')
        self.assertEqual(data['surveyDate'], '2026-09-28')
        self.assertEqual({r['priceJpy'] for r in data['records']}, {181.1, 170.2, 159.3})

    def test_invalid_headers_missing_prefecture_nonnumeric_and_zero(self):
        for cell, bad in [('C5', 'ハイオク（円/18L）'), ('B3', '税抜'), ('B7', None), ('D8', '170.2'), ('F8', 0), ('H8', 401), ('F6', datetime(2026, 9, 14)), ('B8', 'unknown region')]:
            with self.subTest(cell=cell, bad=bad):
                s = self.book[SHEET]
                old = s[cell].value
                s[cell] = bad
                self.book.save(self.path)
                with self.assertRaises(ValueError):
                    self.parse()
                s[cell] = old
        self.book.save(self.path)

    def test_schema_missing_duplicate_nan_and_downgrade(self):
        valid = self.parse()
        for mutate in [lambda d: d.update(unit='JPY/18L'), lambda d: d['records'].pop(), lambda d: d['records'].__setitem__(0, d['records'][1]), lambda d: d['records'][0].update(priceJpy=float('nan'))]:
            data = copy.deepcopy(valid)
            mutate(data)
            with self.assertRaises(ValueError):
                validate_prices(data)
        previous = {'surveyDate': '2026-09-29', 'publishedAt': '2026-09-30'}
        with self.assertRaisesRegex(ValueError, 'Older'):
            validate_prices(valid, previous)

    def test_source_proof_hash_domain_and_dates(self):
        verified_input(self.path, self.meta, 'www.enecho.meti.go.jp')
        for key, value in [('sha256', '0' * 64), ('bytes', 1), ('sourceUrl', 'https://example.com/fake.xlsx')]:
            with self.assertRaises(ValueError):
                verified_input(self.path, self.meta | {key: value}, 'www.enecho.meti.go.jp')
        with self.assertRaises(ValueError):
            parse_prices(self.path, self.meta | {'surveyDate': '2026-09-14'})

    def test_failed_import_keeps_manifest_and_registry(self):
        output = Path(self.temp.name) / 'published'
        shutil.copytree(ROOT / 'public/data', output)
        before = (output / 'manifest.json').read_bytes()
        registry = (output / 'source-registry.json').read_bytes()
        self.book[SHEET]['F8'] = 'invalid'
        self.book.save(self.path)
        proof = Path(self.temp.name) / 'metadata.json'
        proof.write_bytes(encode(metadata(self.path)))
        args = SimpleNamespace(output=str(output), kind='prices', workbook=str(self.path), price_metadata=str(proof))
        with self.assertRaises(ValueError):
            update(args)
        self.assertEqual(before, (output / 'manifest.json').read_bytes())
        self.assertEqual(registry, (output / 'source-registry.json').read_bytes())

    def test_failed_pointer_promotion_restores_registry_and_previous_manifest(self):
        output = Path(self.temp.name) / 'published'
        shutil.copytree(ROOT / 'public/data', output)
        before = (output / 'manifest.json').read_bytes()
        registry = (output / 'source-registry.json').read_bytes()
        # This regression must still reach promotion after a valid weekly refresh.
        latest = json.loads(before)['prices']
        sheet = self.book[SHEET]
        sheet['E2'] = datetime.fromisoformat(latest['publishedAt'])
        survey = datetime.fromisoformat(latest['surveyDate'])
        for cell in ['D6', 'F6', 'H6']:
            sheet[cell] = survey
        for cell in ['C6', 'E6', 'G6']:
            sheet[cell] = survey - timedelta(days=7)
        self.book.save(self.path)
        proof = Path(self.temp.name) / 'metadata.json'
        proof.write_bytes(encode(metadata(self.path)))
        args = SimpleNamespace(output=str(output), kind='prices', workbook=str(self.path), price_metadata=str(proof))
        from common import atomic_write
        def fail_pointer(path, payload):
            if Path(path).name == 'manifest.json':
                raise OSError('Simulated pointer write failure')
            atomic_write(path, payload)
        with patch('pipeline.atomic_write', side_effect=fail_pointer):
            with self.assertRaises(OSError):
                update(args)
        self.assertEqual(before, (output / 'manifest.json').read_bytes())
        self.assertEqual(registry, (output / 'source-registry.json').read_bytes())

    def test_http403_never_replaces_local_or_published_data(self):
        target = Path(self.temp.name) / 'download.xlsx'
        args = SimpleNamespace(url=self.meta['sourceUrl'], destination=str(target))
        error = HTTPError(args.url, 403, 'Forbidden', {}, io.BytesIO())
        self.addCleanup(error.close)
        with patch('urllib.request.urlopen', side_effect=error):
            with self.assertRaises(HTTPError):
                fetch_price(args)
        self.assertFalse(target.exists())
        self.assertFalse(Path(str(target) + '.json').exists())


class OsmTests(unittest.TestCase):
    def test_osm_requires_the_available_publisher_checksum(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / 'checksum-fixture.bin'
            path.write_bytes(b'Fixture only: input checksum validation, not a PBF parser fixture.')
            proof = {'sourceUrl': 'https://download.geofabrik.de/asia/japan-260929.osm.pbf',
                     'fetchedAt': datetime.now(timezone.utc).isoformat(), 'bytes': path.stat().st_size,
                     'sha256': digest(path), 'md5': digest(path, 'md5')}
            verified_input(path, proof, 'download.geofabrik.de')
            for bad in [None, '', 'bad', 'g' * 32, '0' * 32]:
                changed = dict(proof)
                if bad is None:
                    changed.pop('md5')
                else:
                    changed['md5'] = bad
                with self.subTest(md5=bad), self.assertRaises(ValueError):
                    verified_input(path, changed, 'download.geofabrik.de')
            verified_input(path, proof | {'md5': proof['md5'].upper()}, 'download.geofabrik.de')

    def test_unknown_mapping_and_conflicting_evidence(self):
        self.assertEqual(tri({}, ['fuel:regular']), 'UNKNOWN')
        self.assertEqual(tri({'fuel:gasoline': 'yes'}, ['fuel:regular']), 'UNKNOWN')
        self.assertEqual(tri({'fuel:octane_91': 'yes'}, ['fuel:regular']), 'UNKNOWN')
        self.assertEqual(tri({'fuel:regular': 'yes', 'fuel:レギュラー': 'no'}, ['fuel:regular', 'fuel:レギュラー']), 'UNKNOWN')
        self.assertEqual(tri({'payment:visa': 'yes', 'payment:credit_cards': 'no'}, ['payment:visa'], 'payment:credit_cards'), 'UNKNOWN')
        self.assertEqual(tri({'fuel:diesel': 'no'}, ['fuel:diesel']), 'NO')

    def test_conditional_attributes_remain_unknown(self):
        for tags, keys, parent in [
            ({'payment:visa': 'yes', 'payment:visa:conditional': 'no @ (night)'}, ['payment:visa'], 'payment:credit_cards'),
            ({'fuel:diesel': 'yes', 'fuel:diesel:conditional': 'no @ (night)'}, ['fuel:diesel'], None),
            ({'payment:visa': 'yes', 'payment:credit_cards:conditional': 'no @ (night)'}, ['payment:visa'], 'payment:credit_cards'),
        ]:
            with self.subTest(tags=tags):
                self.assertEqual(tri(tags, keys, parent), 'UNKNOWN')

    def test_service_never_collapses_mixed_or_conditional_evidence(self):
        from shapely.geometry import Point
        base = {'@type': 'node', '@id': 1, '@timestamp': 1700000000, 'amenity': 'fuel'}
        cases = [
            ({'self_service': 'yes'}, 'SELF'),
            ({'self_service': 'only'}, 'SELF'),
            ({'self_service': 'no'}, 'UNKNOWN'),
            ({'self_service': 'no', 'full_service': 'yes'}, 'FULL'),
            ({'self_service': 'yes', 'full_service': 'yes'}, 'UNKNOWN'),
            ({'self_service': 'yes', 'self_service:conditional': 'no @ (06:00-18:00)'}, 'UNKNOWN'),
            ({'self_service': 'yes', 'opening_hours:self_service': 'Mo-Fr'}, 'UNKNOWN'),
            ({'self_service': 'no', 'full_service': 'yes', 'full_service:conditional': 'no @ (night)'}, 'UNKNOWN'),
        ]
        for tags, expected in cases:
            with self.subTest(tags=tags):
                result = station({'properties': base | tags}, Point(139, 35), 'JP-13', 'OSM_NODE')
                self.assertEqual(result['serviceType'], expected)

    def test_count_gate_bootstrap_all_partitions_and_manual_reason(self):
        old = dict.fromkeys(CODES + ['UNKNOWN'], 100)
        new = old | {'JP-01': 126}
        self.assertEqual(count_gate(None, old, True)['kind'], 'BOOTSTRAP')
        for arguments in [(None, old), (old, new), (old, new, True), (old, new, False, 'ok')]:
            with self.assertRaises(ValueError):
                count_gate(*arguments)
        self.assertEqual(count_gate(old, new, review_reason='Manually reviewed upstream coverage changes')['kind'], 'MANUAL_REVIEW')
        self.assertEqual(count_gate(old, old | {'JP-01': 125})['kind'], 'WITHIN_THRESHOLD')
        with self.assertRaises(ValueError):
            count_gate(old | {'UNKNOWN': 0}, old | {'UNKNOWN': 1})

    def test_typed_dedup_surface_point_unknown_and_repairs(self):
        features = []
        for i, code in enumerate(CODES):
            x = 123 + i * .2
            features.append({'type': 'Feature', 'properties': {'@type': 'relation', '@id': i + 1, 'boundary': 'administrative', 'admin_level': '4', 'ISO3166-2': code}, 'geometry': {'type': 'Polygon', 'coordinates': [[[x, 30], [x + .15, 30], [x + .15, 31], [x, 31], [x, 30]]]}})
        tags = {'@type': 'way', '@id': 1, '@timestamp': 1700000000, 'amenity': 'fuel', 'self_service': 'no', 'brand': 'ENEOS', 'fuel:gasoline': 'yes'}
        # Bowtie area is repaired explicitly; line duplicate must be discarded by type rank.
        ring = [[123.01, 30.1], [123.1, 30.2], [123.01, 30.2], [123.1, 30.1], [123.01, 30.1]]
        features.extend([{'type': 'Feature', 'properties': tags, 'geometry': {'type': 'LineString', 'coordinates': ring}}, {'type': 'Feature', 'properties': tags, 'geometry': {'type': 'Polygon', 'coordinates': [ring]}}, {'type': 'Feature', 'properties': tags | {'@type': 'node'}, 'geometry': {'type': 'Point', 'coordinates': [140, 40]}}])
        with tempfile.NamedTemporaryFile(mode='w') as file:
            file.write('\n'.join(json.dumps(f) for f in features));file.flush()
            partitions, audit = parse_osm(file.name)
        self.assertEqual(audit['inputFuelFeatures'], 3)
        self.assertEqual(audit['uniqueObjects'], 2)
        self.assertEqual(audit['duplicatesRemoved'], 1)
        self.assertEqual(len(audit['repaired']), 1)
        self.assertEqual(partitions['JP-01'][0]['positionMethod'], 'SURFACE_POINT')
        self.assertEqual(partitions['JP-01'][0]['serviceType'], 'UNKNOWN')
        self.assertEqual(partitions['JP-01'][0]['fuelRegular'], 'UNKNOWN')
        self.assertEqual(partitions['UNKNOWN'][0]['id'], 'osm:node:1')
        self.assertEqual(audit['unknownAssignments'][0]['reason'], 'OUTSIDE_AVAILABLE_BOUNDARIES')

if __name__ == '__main__':
    unittest.main()
